from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ARCHIVO_ORIGINAL = Path("marcacionesSantani3.xlsx")
ARCHIVO_SALIDA = Path("marcacionesSantani3_con_reporte.xlsx")
NOMBRE_HOJA_REPORTE = "Reporte_CI_Fecha"


def convertir_fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if hasattr(valor, "year") and hasattr(valor, "month") and hasattr(valor, "day"):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    return None


def convertir_hora(valor):
    if hasattr(valor, "hour") and hasattr(valor, "minute") and hasattr(valor, "second"):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(texto, formato).time()
        except ValueError:
            pass
    return None


def main():
    if not ARCHIVO_ORIGINAL.exists():
        raise FileNotFoundError(f"No encontré el archivo: {ARCHIVO_ORIGINAL}")

    libro = load_workbook(ARCHIVO_ORIGINAL)
    hoja_origen = libro.active

    encabezados = {
        str(celda.value).strip().upper(): celda.column
        for celda in hoja_origen[1]
        if celda.value is not None
    }
    requeridas = {"CI", "FECHA", "HORA"}
    faltantes = requeridas - set(encabezados)
    if faltantes:
        raise ValueError(f"Faltan columnas requeridas: {sorted(faltantes)}")

    agrupado = defaultdict(list)
    filas_invalidas = 0

    for fila in hoja_origen.iter_rows(min_row=2, values_only=True):
        ci = str(fila[encabezados["CI"] - 1] or "").strip()
        fecha = convertir_fecha(fila[encabezados["FECHA"] - 1])
        hora = convertir_hora(fila[encabezados["HORA"] - 1])

        if not ci or fecha is None or hora is None:
            filas_invalidas += 1
            continue

        agrupado[(ci, fecha)].append(datetime.combine(fecha, hora))

    if NOMBRE_HOJA_REPORTE in libro.sheetnames:
        del libro[NOMBRE_HOJA_REPORTE]
    hoja_reporte = libro.create_sheet(NOMBRE_HOJA_REPORTE)

    columnas = [
        "CI",
        "FECHA",
        "PRIMERA_MARCACION",
        "ULTIMA_MARCACION",
        "TOTAL_MARCACIONES",
    ]
    hoja_reporte.append(columnas)

    for celda in hoja_reporte[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")

    for (ci, fecha), marcas in sorted(agrupado.items(), key=lambda item: (item[0][0], item[0][1])):
        marcas.sort()
        hoja_reporte.append([
            ci,
            fecha,
            marcas[0].time(),
            marcas[-1].time(),
            len(marcas),
        ])

    for fila in hoja_reporte.iter_rows(min_row=2):
        fila[1].number_format = "dd/mm/yyyy"
        fila[2].number_format = "hh:mm:ss"
        fila[3].number_format = "hh:mm:ss"

    hoja_reporte.freeze_panes = "A2"
    hoja_reporte.auto_filter.ref = hoja_reporte.dimensions
    anchos = [16, 14, 22, 22, 22]
    for indice, ancho in enumerate(anchos, start=1):
        hoja_reporte.column_dimensions[get_column_letter(indice)].width = ancho

    libro.save(ARCHIVO_SALIDA)
    total_marcaciones = sum(len(marcas) for marcas in agrupado.values())
    print(f"Archivo generado: {ARCHIVO_SALIDA}")
    print(f"Marcaciones procesadas: {total_marcaciones}")
    print(f"Filas omitidas por datos inválidos: {filas_invalidas}")


if __name__ == "__main__":
    main()
